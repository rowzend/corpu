from rest_framework import generics, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.core.management import call_command
import math
from django.db.models import Q
from django.utils import timezone
from io import StringIO
import json
import openpyxl

from .models import MsPerguruanTinggi, MsProgramStudi, MsInstansi
from .models import MsProvinsi, MsKabupaten, MsKecamatan, MsKelurahan, MsKategoriUser
from .serializers import (
    PerguruanTinggiSerializer, ProgramStudiSerializer, InstansiSerializer,
    ProvinsiSerializer, KabupatenSerializer, KecamatanSerializer, KelurahanSerializer,
    KategoriUserSerializer,
)


class PerguruanTinggiListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = PerguruanTinggiSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsPerguruanTinggi.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        if search:
            qs = qs.filter(
                Q(nama_pt__icontains=search) |
                Q(kode_pt__icontains=search) |
                Q(kota__icontains=search)
            )
        return qs.order_by('nama_pt')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True,
            'data': serializer.data,
            'pagination': {
                'page': page,
                'page_size': page_size,
                'total': total,
                'total_pages': total_pages,
            }
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class PerguruanTinggiDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = PerguruanTinggiSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsPerguruanTinggi.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


class ProgramStudiListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ProgramStudiSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsProgramStudi.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        pt_id = self.request.query_params.get('perguruan_tinggi_id', '')
        if search:
            qs = qs.filter(
                Q(nama_prodi__icontains=search) |
                Q(kode_prodi__icontains=search) |
                Q(perguruan_tinggi__nama_pt__icontains=search)
            )
        if pt_id:
            qs = qs.filter(perguruan_tinggi_id=pt_id)
        return qs.order_by('perguruan_tinggi__nama_pt', 'jenjang', 'nama_prodi')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True,
            'data': serializer.data,
            'pagination': {
                'page': page,
                'page_size': page_size,
                'total': total,
                'total_pages': total_pages,
            }
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class ProgramStudiDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ProgramStudiSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsProgramStudi.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


class InstansiListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = InstansiSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsInstansi.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        if search:
            qs = qs.filter(
                Q(nama_instansi__icontains=search) |
                Q(kode_instansi__icontains=search) |
                Q(jenis_instansi__icontains=search)
            )
        return qs.order_by('nama_instansi')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True,
            'data': serializer.data,
            'pagination': {
                'page': page,
                'page_size': page_size,
                'total': total,
                'total_pages': total_pages,
            }
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class InstansiDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = InstansiSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsInstansi.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


# ==========================================
# Wilayah Views (Provinsi, Kabupaten, Kecamatan, Kelurahan)
# ==========================================

class ProvinsiListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ProvinsiSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsProvinsi.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        if search:
            qs = qs.filter(Q(nama__icontains=search) | Q(kode__icontains=search))
        return qs.order_by('kode')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        all_val = request.query_params.get('all', '')
        if all_val == 'true':
            serializer = self.get_serializer(queryset, many=True)
            return Response({'success': True, 'data': serializer.data})
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True, 'data': serializer.data,
            'pagination': {'page': page, 'page_size': page_size, 'total': total, 'total_pages': total_pages},
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class ProvinsiDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ProvinsiSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsProvinsi.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


class KabupatenListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KabupatenSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsKabupaten.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        provinsi_id = self.request.query_params.get('provinsi_id', '')
        if search:
            qs = qs.filter(Q(nama__icontains=search) | Q(kode__icontains=search))
        if provinsi_id:
            qs = qs.filter(provinsi_id=provinsi_id)
        return qs.order_by('kode')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        all_val = request.query_params.get('all', '')
        if all_val == 'true':
            serializer = self.get_serializer(queryset, many=True)
            return Response({'success': True, 'data': serializer.data})
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True, 'data': serializer.data,
            'pagination': {'page': page, 'page_size': page_size, 'total': total, 'total_pages': total_pages},
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class KabupatenDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KabupatenSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsKabupaten.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


class KecamatanListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KecamatanSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsKecamatan.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        kabupaten_id = self.request.query_params.get('kabupaten_id', '')
        if search:
            qs = qs.filter(Q(nama__icontains=search) | Q(kode__icontains=search))
        if kabupaten_id:
            qs = qs.filter(kabupaten_id=kabupaten_id)
        return qs.order_by('kode')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        all_val = request.query_params.get('all', '')
        if all_val == 'true':
            serializer = self.get_serializer(queryset, many=True)
            return Response({'success': True, 'data': serializer.data})
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True, 'data': serializer.data,
            'pagination': {'page': page, 'page_size': page_size, 'total': total, 'total_pages': total_pages},
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class KecamatanDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KecamatanSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsKecamatan.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


class KelurahanListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KelurahanSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsKelurahan.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        kecamatan_id = self.request.query_params.get('kecamatan_id', '')
        if search:
            qs = qs.filter(Q(nama__icontains=search) | Q(kode__icontains=search))
        if kecamatan_id:
            qs = qs.filter(kecamatan_id=kecamatan_id)
        return qs.order_by('kode')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        all_val = request.query_params.get('all', '')
        if all_val == 'true':
            serializer = self.get_serializer(queryset, many=True)
            return Response({'success': True, 'data': serializer.data})
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True, 'data': serializer.data,
            'pagination': {'page': page, 'page_size': page_size, 'total': total, 'total_pages': total_pages},
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class KelurahanDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KelurahanSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsKelurahan.objects.filter(deleted_at__isnull=True)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        serializer = self.get_serializer(instance)
        return Response({'success': True, 'data': serializer.data})

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response({'success': True, 'data': serializer.data})

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        instance.delete()
        return Response({'success': True, 'message': 'Berhasil dihapus'}, status=status.HTTP_200_OK)


JENIS_INSTANSI_MAP = {
    'KO': 'Kementerian Koordinator',
    'KEMENT': 'Kementerian',
    'LPNK': 'Lembaga Non Kementerian',
    'LNS': 'Lembaga Non Struktural',
    'PROV': 'Provinsi',
    'KAB': 'Kabupaten',
    'KOT': 'Kota',
}

TINGKAT_INSTANSI_MAP = {
    'P': 'Pusat',
    'D': 'Daerah',
}


class InstansiImportXLSXAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'success': False, 'message': 'File tidak ditemukan'}, status=status.HTTP_400_BAD_REQUEST)

        if not file.name.endswith(('.xlsx', '.xls')):
            return Response({'success': False, 'message': 'Format file harus .xlsx atau .xls'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(file, read_only=True)
            ws = wb.active

            first_cell = ''
            for c in next(ws.iter_rows(min_row=1, max_row=1)):
                if c.value:
                    first_cell = str(c.value).strip().upper()
                    break

            if first_cell == 'TABREF INSTANSI_ID':
                total = self._import_bkn_format(ws)
            else:
                total = self._import_custom_format(ws)

            return Response({
                'success': True,
                'message': f'Import selesai: {total} instansi tersimpan'
            })
        except Exception as e:
            return Response({
                'success': False,
                'message': f'Gagal membaca file: {str(e)}'
            }, status=status.HTTP_400_BAD_REQUEST)

    def _import_bkn_format(self, ws):
        total = 0
        for row in ws.iter_rows(min_row=4, values_only=True):
            if len(row) < 5:
                continue
            id_bkn = str(row[0]).strip() if row[0] else ''
            nama = str(row[1]).strip() if row[1] else ''
            jenis = str(row[2]).strip().upper() if row[2] else ''
            cepat_kode = str(row[3]).strip() if row[3] else ''
            jenis_id = str(row[4]).strip().upper() if row[4] else ''

            if not nama:
                continue
            if not cepat_kode:
                continue

            try:
                MsInstansi.objects.update_or_create(
                    kode_instansi=cepat_kode,
                    defaults={
                        'id_bkn': id_bkn or None,
                        'nama_instansi': nama,
                        'jenis_instansi': JENIS_INSTANSI_MAP.get(jenis_id, jenis_id or None),
                        'tingkat_instansi': TINGKAT_INSTANSI_MAP.get(jenis, jenis or None),
                        'is_active': True,
                    }
                )
                total += 1
            except Exception:
                pass
        return total

    def _import_custom_format(self, ws):
        headers = [str(c.value).strip().lower() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {}
        for h in ['kode_instansi', 'nama_instansi', 'jenis_instansi', 'tingkat_instansi', 'alamat', 'telepon', 'website', 'email', 'id_bkn']:
            for i, col in enumerate(headers):
                if col.replace(' ', '_') == h or col == h:
                    col_map[h] = i
                    break

        if 'kode_instansi' not in col_map or 'nama_instansi' not in col_map:
            raise Exception('Header wajib: kode_instansi, nama_instansi')

        total = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[col_map['kode_instansi']] or not row[col_map['nama_instansi']]:
                continue
            try:
                data = {}
                for field, idx in col_map.items():
                    val = str(row[idx]).strip() if idx < len(row) and row[idx] else ''
                    data[field] = val if val else None
                MsInstansi.objects.update_or_create(
                    kode_instansi=data['kode_instansi'],
                    defaults={
                        'id_bkn': data.get('id_bkn'),
                        'nama_instansi': data['nama_instansi'],
                        'jenis_instansi': data.get('jenis_instansi'),
                        'tingkat_instansi': data.get('tingkat_instansi'),
                        'alamat': data.get('alamat'),
                        'telepon': data.get('telepon'),
                        'website': data.get('website'),
                        'email': data.get('email'),
                        'is_active': True,
                    }
                )
                total += 1
            except Exception:
                pass
        return total


class KategoriUserListAPIView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KategoriUserSerializer
    pagination_class = None

    def get_queryset(self):
        qs = MsKategoriUser.objects.filter(deleted_at__isnull=True)
        search = self.request.query_params.get('search', '')
        if search:
            qs = qs.filter(Q(nama__icontains=search) | Q(kode__icontains=search))
        return qs.order_by('kode')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        all_val = request.query_params.get('all', '')
        if all_val == 'true':
            serializer = self.get_serializer(queryset, many=True)
            return Response({'success': True, 'data': serializer.data})
        page_size = int(request.query_params.get('page_size', 20))
        page = int(request.query_params.get('page', 1))
        total = queryset.count()
        total_pages = max(math.ceil(total / page_size) if page_size > 0 else 0, 1)
        start = (page - 1) * page_size
        end = start + page_size
        serializer = self.get_serializer(queryset[start:end], many=True)
        return Response({
            'success': True, 'data': serializer.data,
            'pagination': {'page': page, 'page_size': page_size, 'total': total, 'total_pages': total_pages},
        })

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response({'success': True, 'data': serializer.data}, status=status.HTTP_201_CREATED)


class KategoriUserDetailAPIView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = KategoriUserSerializer
    lookup_field = 'pk'

    def get_queryset(self):
        return MsKategoriUser.objects.filter(deleted_at__isnull=True)

    def perform_destroy(self, instance):
        instance.deleted_at = timezone.now()
        instance.save()


class SyncReferensiAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        source = request.data.get('source', 'github')
        try:
            out = StringIO()
            if source == 'pddikti':
                try:
                    call_command('seed_referensi_data', '--source=pddikti', stdout=out)
                except Exception:
                    out.write('PD-DIKTI tidak dapat dijangkau. Fallback ke GitHub.\n')
                    call_command('seed_referensi_data', '--source=github', stdout=out)
            elif source == 'github':
                call_command('seed_referensi_data', '--source=github', stdout=out)
            elif source == 'xlsx':
                call_command('seed_referensi_data', '--source=xlsx', stdout=out)
            elif source == 'sample':
                call_command('seed_referensi_data', '--source=sample', stdout=out)
            elif source == 'wilayah':
                call_command('seed_wilayah_data', stdout=out)
            else:
                call_command('seed_referensi_data', f'--source={source}', stdout=out)
            output = out.getvalue()
            return Response({'success': True, 'message': 'Sinkronisasi berhasil', 'output': output})
        except Exception as e:
            return Response({
                'success': False,
                'message': f'Sinkronisasi gagal: {str(e)}'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
